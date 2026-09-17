#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ПРЕДВАРИТЕЛЬНЫЙ ЗАКАЗ НОВОГОДНЕГО КОРПОРАТИВА
=============================================
Генератор премиальной презентации (8 слайдов, 16:9) для ивент-агентства.

Стиль:   dark luxury / Black Tie — глубокий midnight blue + champagne gold,
         тонкие золотые линии, матовые «стеклянные» плашки, много воздуха.
Шрифты:  Montserrat (заголовки, нумерация) + Inter (текст) — free, Google Fonts.
         Если шрифтов нет — PowerPoint корректно подставит fallback.
Решётка: единая модульная сетка 12 колонок, левая колонна (текст) 6.9",
         правая (визуал) 4.28", поля 0.9" по горизонтали.

Запуск:  python build_presentation.py
Результат:
  preorder_new_year_corporate_2025.pptx  — сама презентация
  scene.json                             — манифест геометрии (для превью/ревью)

Требования: pip install python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from lxml import etree
import json
import os

# ---------------------------------------------------------------------------
# ПАЛИТРА И КОНСТАНТЫ
# ---------------------------------------------------------------------------
MIDNIGHT   = RGBColor(0x0A, 0x0F, 0x1D)   # базовый ночной фон
GOLD       = RGBColor(0xD4, 0xAF, 0x37)   # champagne gold
GOLD_LIGHT = RGBColor(0xE5, 0xC3, 0x78)   # морозное золото (акценты текста)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GRAY       = RGBColor(0xA0, 0xAE, 0xC0)   # второстепенный текст
INK        = RGBColor(0x0A, 0x0F, 0x1D)   # тёмный текст на золоте

F_DISPLAY = "Montserrat"   # заголовки, нумерация, бейджи
F_BODY    = "Inter"        # основной текст

# Габариты слайда 16:9
W, H = 13.333, 7.5
ML = MR = 0.9                       # поля
LX, LW = 0.9, 6.9                   # левая колонна (текст)
RX, RW = 8.15, 4.28                 # правая колонна (визуал)
RXI, RWI = 8.47, 3.64               # внутренние отступы правой карточки
HDR_Y, FOOT_Y = 1.12, 7.02          # линии шапки / подвала

# ---------------------------------------------------------------------------
# СЦЕНА (манифест геометрии для рендера превью)
# ---------------------------------------------------------------------------
SCENE = {"w": W, "h": H, "slides": []}


def _sl():
    return SCENE["slides"][-1]


def _rgb(c):
    return (c[0], c[1], c[2])


# ---------------------------------------------------------------------------
# БАЗОВЫЕ ХЕЛПЕРЫ
# ---------------------------------------------------------------------------
def _add_alpha(color_el, pct):
    """Добавить прозрачность (0..100) к элементу a:srgbClr."""
    if color_el is None:
        return
    for old in color_el.findall(qn("a:alpha")):
        color_el.remove(old)
    el = etree.SubElement(color_el, qn("a:alpha"))
    el.set("val", str(int(round(pct * 1000))))


def _shape_fill_alpha(sp, pct):
    sf = sp._element.spPr.find(qn("a:solidFill"))
    if sf is not None:
        _add_alpha(sf.find(qn("a:srgbClr")), pct)


def _shape_line_alpha(sp, pct):
    ln = sp._element.spPr.find(qn("a:ln"))
    if ln is not None:
        sf = ln.find(qn("a:solidFill"))
        if sf is not None:
            _add_alpha(sf.find(qn("a:srgbClr")), pct)


def _set_gradient(sp, stops):
    """Вертикальный градиент: stops = [(pos 0..100, 'RRGGBB'), ...]."""
    spPr = sp._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        e = spPr.find(qn(tag))
        if e is not None:
            spPr.remove(e)
    gs = "".join(
        '<a:gs pos="%d"><a:srgbClr val="%s"/></a:gs>' % (int(p * 1000), h)
        for p, h in stops
    )
    xml = ('<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
           ' rotWithShape="1"><a:gsLst>%s</a:gsLst>'
           '<a:lin ang="5400000" scaled="0"/></a:gradFill>') % gs
    grad = parse_xml(xml)
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(grad)
    else:
        spPr.append(grad)


def box(slide, st, x, y, w, h, fill=None, fill_a=None, line=None, line_a=None,
        line_w=0.75, dash=None, radius=None, rot=None, grad=None):
    """Прямоугольник/овал с точным контролем заливки, рамки и прозрачности."""
    sp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if grad is None:
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = fill
            if fill_a is not None:
                _shape_fill_alpha(sp, fill_a)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
        if line_a is not None:
            _shape_line_alpha(sp, line_a)
        if dash:
            sp.line.dash_style = dash
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if rot:
        sp.rotation = rot
    if grad is not None:
        _set_gradient(sp, grad)
    _sl()["shapes"].append({
        "kind": "roundrect" if radius is not None else
                ("ellipse" if st == MSO_SHAPE.OVAL else
                 ("diamond" if st == MSO_SHAPE.DIAMOND else "rect")),
        "x": x, "y": y, "w": w, "h": h, "rot": rot or 0,
        "fill": (_rgb(fill), fill_a if fill_a is not None else 100) if fill is not None and grad is None else None,
        "line": (_rgb(line), line_a if line_a is not None else 100, line_w, dash.name if dash else None) if line is not None else None,
        "radius": radius, "grad": grad,
    })
    return sp


def P(t, s, c=WHITE, f=F_BODY, b=False, spc=0, a=100):
    """Один run: текст t, размер s pt, цвет c, шрифт f, bold, трекинг spc, альфа a."""
    return {"t": t, "px": s, "c": [c[0], c[1], c[2], a], "f": f, "b": b, "spc": spc, "a": a}


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    """Многострочный текст. paras = [ {runs, align, ls, before, after} ... ]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    log_paras = []
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        if p.get("before"):
            para.space_before = Pt(p["before"])
        if p.get("after"):
            para.space_after = Pt(p["after"])
        log_runs = []
        for r in p["runs"]:
            run = para.add_run()
            run.text = r["t"]
            f = run.font
            f.size = Pt(r["px"])
            f.name = r["f"]
            f.bold = bool(r.get("b"))
            f.color.rgb = RGBColor(r["c"][0], r["c"][1], r["c"][2])
            rPr = run._r.get_or_add_rPr()
            rPr.set("spc", str(int(r.get("spc", 0) * 100)))
            if r.get("a") is not None and r["a"] != 100:
                sf = rPr.find(qn("a:solidFill"))
                if sf is not None:
                    _add_alpha(sf.find(qn("a:srgbClr")), r["a"])
            log_runs.append(r)
        log_paras.append({
            "runs": log_runs,
            "align": {PP_ALIGN.LEFT: "left", PP_ALIGN.CENTER: "center",
                      PP_ALIGN.RIGHT: "right"}[p.get("align", PP_ALIGN.LEFT)],
            "ls": p.get("ls"), "before": p.get("before"), "after": p.get("after"),
        })
    _sl()["texts"].append({"x": x, "y": y, "w": w, "h": h,
                           "anchor": "middle" if anchor == MSO_ANCHOR.MIDDLE else "top",
                           "wrap": wrap, "paras": log_paras})
    return tb


def line_h(slide, x, y, w, c, a, wpt=0.75, dash=None):
    """Горизонтальная тонкая линия (hairline)."""
    return box(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.011, fill=c, fill_a=a, dash=dash)


def line_v(slide, x, y, h, c, a, wpt=0.75):
    """Вертикальная тонкая линия."""
    return box(slide, MSO_SHAPE.RECTANGLE, x, y, 0.011, h, fill=c, fill_a=a)


def diamond(slide, cx, cy, s, c=GOLD, a=100):
    """Маленький золотой ромб (декоративный маркер)."""
    return box(slide, MSO_SHAPE.DIAMOND, cx - s / 2, cy - s / 2, s, s, fill=c, fill_a=a)


LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         "assets", "logo_white.png")


def logo(slide, x, y, h):
    """Знак «Кьево» (белый, прозрачный фон). Квадратная форма: h — высота."""
    pic = slide.shapes.add_picture(LOGO_PATH, Inches(x), Inches(y), Inches(h), Inches(h))
    pic.shadow.inherit = False
    _sl()["pics"].append({"x": x, "y": y, "w": h, "h": h,
                          "src": os.path.relpath(LOGO_PATH,
                                                 os.path.dirname(os.path.abspath(__file__)))})
    return pic


# ---------------------------------------------------------------------------
# СОВМЕСТНЫЕ ЭЛЕМЕНТЫ СЛАЙДОВ 2–8 (единый брендинг)
# ---------------------------------------------------------------------------
def base_slide(prs, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = MIDNIGHT
    SCENE["slides"].append({"idx": idx, "shapes": [], "texts": [], "pics": []})
    # Фоновый вертикальный градиент (глубина, без «клипарта»)
    box(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, H,
        grad=[(0, "111A30"), (55, "0B101F"), (100, "080C16")])
    # Шапка: оверлей-метка слева, логотип справа, золотая линия
    text(slide, 0.9, 0.52, 5.5, 0.3,
         [{"runs": [P("МОСКВА · СЕЗОН 2024–2025", 8, GRAY, F_DISPLAY, spc=350, a=70)]}])
    logo(slide, 12.03, 0.42, 0.4)
    line_h(slide, 0.9, HDR_Y, W - ML - MR, GOLD, 22)
    # Подвал: линия, номер слайда, подпись бренда
    line_h(slide, 0.9, FOOT_Y, W - ML - MR, WHITE, 10)
    text(slide, 0.9, 7.14, 2.5, 0.28,
         [{"runs": [P("%02d / 08" % idx, 8, GRAY, F_BODY, spc=200, a=60)]}])
    text(slide, 7.43, 7.14, 5.0, 0.28,
         [{"runs": [P("НОВОГОДНИЙ КОРПОРАТИВ · ПРЕДВАРИТЕЛЬНЫЙ ЗАКАЗ", 8, GRAY,
                      F_BODY, spc=200, a=45)], "align": PP_ALIGN.RIGHT}])
    return slide


def ghost_number(slide, num):
    """Крупная полупрозрачная цифра преимущества на фоне."""
    text(slide, 7.3, 1.85, 5.13, 3.6,
         [{"runs": [P(num, 215, GOLD, F_DISPLAY, b=True, a=8)],
           "align": PP_ALIGN.RIGHT}], wrap=False)


def kicker(slide, label):
    """Золотой ромб + надзаголовок преимущества."""
    diamond(slide, 0.945, 1.685, 0.09)
    text(slide, 1.12, 1.55, 8.0, 0.32,
         [{"runs": [P(label, 10, GOLD_LIGHT, F_DISPLAY, b=True, spc=350)]}])


def heading(slide, title):
    text(slide, LX, 1.98, LW + 0.1, 1.5,
         [{"runs": [P(title, 28, WHITE, F_DISPLAY, b=True, spc=50)], "ls": 1.12}])


def body(slide, txt):
    text(slide, LX, 3.62, 6.75, 2.8,
         [{"runs": [P(txt, 12.5, GRAY, F_BODY)], "ls": 1.42}])


# ---------------------------------------------------------------------------
# СЛАЙД 1 — ОБЛОЖКА
# ---------------------------------------------------------------------------
def slide_01(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = MIDNIGHT
    SCENE["slides"].append({"idx": 1, "shapes": [], "texts": [], "pics": []})
    box(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, H,
        grad=[(0, "111A30"), (55, "0B101F"), (100, "080C16")])

    # Декор: два тонких кольца (глубина, статус), по краям — уголки-рамки
    box(slide, MSO_SHAPE.OVAL, 11.35, -1.45, 3.9, 3.9, line=GOLD, line_a=14, line_w=1.0)
    box(slide, MSO_SHAPE.OVAL, -1.15, 5.75, 2.6, 2.6, line=WHITE, line_a=10, line_w=1.0)
    for x, y, flip_x, flip_y in [
        (0.5, 0.5, 1, 1), (12.833, 0.5, -1, 1), (0.5, 7.0, 1, -1), (12.833, 7.0, -1, -1)
    ]:
        x0 = x if flip_x > 0 else x - 0.22
        y0 = y if flip_y > 0 else y - 0.22
        box(slide, MSO_SHAPE.RECTANGLE, x0, y if flip_y > 0 else y, 0.22, 0.011,
            fill=GOLD, fill_a=35)
        box(slide, MSO_SHAPE.RECTANGLE, x, y0, 0.011, 0.22, fill=GOLD, fill_a=35)

    # Логотип — крупно по центру сверху
    logo(slide, (W - 0.95) / 2, 0.68, 0.95)

    # Центральный блок
    diamond(slide, W / 2, 2.22, 0.09, GOLD, 90)
    text(slide, 1.42, 2.44, 10.5, 0.32,
         [{"runs": [P("EXCLUSIVE PRE-BOOKING · WINTER 2024–2025", 9, GRAY,
                      F_DISPLAY, b=True, spc=500, a=75)], "align": PP_ALIGN.CENTER}])
    text(slide, 1.42, 2.72, 10.5, 1.7,
         [{"runs": [P("ПРЕДВАРИТЕЛЬНЫЙ ЗАКАЗ", 36, WHITE, F_DISPLAY, b=True, spc=300)],
           "align": PP_ALIGN.CENTER, "ls": 1.18},
          {"runs": [P("НОВОГОДНЕГО КОРПОРАТИВА", 36, WHITE, F_DISPLAY, b=True, spc=300)],
           "align": PP_ALIGN.CENTER, "ls": 1.18}])
    text(slide, 1.42, 4.44, 10.5, 0.5,
         [{"runs": [P("7 причин зафиксировать идеальный финал года уже сейчас",
                      17, GOLD_LIGHT, F_BODY, b=True, spc=120)],
           "align": PP_ALIGN.CENTER}])
    box(slide, MSO_SHAPE.RECTANGLE, W / 2 - 0.45, 5.2, 0.9, 0.012, fill=GOLD, fill_a=80)

    # Нижний тег
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.37, 6.14, 10.6, 0.56,
        fill=WHITE, fill_a=4, line=WHITE, line_a=16, line_w=0.75, radius=0.5)
    text(slide, 1.37, 6.14, 10.6, 0.56,
         [{"runs": [P("МОСКВА   ·   СЕЗОН 2024–2025   ·   ЭКСКЛЮЗИВНОЕ ПРЕДЛОЖЕНИЕ "
                      "ДЛЯ КОРПОРАТИВНЫХ КЛИЕНТОВ", 9, GRAY, F_BODY, spc=250, a=90)],
           "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# СЛАЙД 2 — ПЛОЩАДКИ И ДАТЫ
# ---------------------------------------------------------------------------
def slide_02(prs):
    slide = base_slide(prs, 2)
    ghost_number(slide, "01")
    kicker(slide, "ПРЕИМУЩЕСТВО 01 · ДОСТУП К ЭКСКЛЮЗИВУ")
    heading(slide, "Лучшие площадки и даты")
    body(slide, "Успеваем перехватить «золотые даты» декабря и бронируем топовые "
                "залы столицы, пока весь бизнес стоит в листах ожидания. Никаких "
                "компромиссов: ваш корпоратив пройдет в идеальный вечер и на самой "
                "эффектной площадке Москвы.")

    # Карточка-акцент «Прайм-тайм декабря гарантирован»
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, RX, 1.98, RW, 4.1,
        fill=WHITE, fill_a=4, line=WHITE, line_a=12, line_w=0.75, radius=0.045)
    box(slide, MSO_SHAPE.RECTANGLE, 8.45, 2.0, 3.68, 0.014, fill=GOLD, fill_a=80)
    text(slide, RXI, 2.34, RWI, 0.3,
         [{"runs": [P("ДЕКАБРЬ · ЗОЛОТЫЕ ДАТЫ", 8.5, GRAY, F_DISPLAY, b=True, spc=300, a=80)]}])
    dates = [["18.12", "19.12", "20.12", "21.12"],
             ["24.12", "25.12", "26.12", "27.12"]]
    for r, row in enumerate(dates):
        for i, d in enumerate(row):
            x = 8.47 + i * (0.83 + 0.066)
            y = 2.72 + r * 0.52
            hot = (d == "27.12")
            box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 0.83, 0.42,
                fill=GOLD if hot else WHITE,
                fill_a=100 if hot else 5,
                line=GOLD if hot else WHITE, line_a=100 if hot else 14, line_w=0.75,
                radius=0.5)
            text(slide, x, y, 0.83, 0.42,
                 [{"runs": [P(d, 10, INK if hot else WHITE, F_BODY, b=hot, spc=50,
                              a=100 if hot else 95)], "align": PP_ALIGN.CENTER}],
                 anchor=MSO_ANCHOR.MIDDLE)
    text(slide, RXI, 3.94, RWI, 0.55,
         [{"runs": [P("Топовые залы, лофты-иконки и террасы столицы — "
                      "бронь уже в работе", 9, GRAY, F_BODY, a=85)], "ls": 1.35}])
    # Плашка-гарантия
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, RXI, 4.74, RWI, 0.74,
        fill=GOLD, fill_a=7, line=GOLD, line_a=85, line_w=1.0, radius=0.10)
    text(slide, RXI, 4.74, RWI, 0.74,
         [{"runs": [P("ПРАЙМ-ТАЙМ ДЕКАБРЯ", 11, GOLD_LIGHT, F_DISPLAY, b=True, spc=300)],
           "align": PP_ALIGN.CENTER, "ls": 1.3},
          {"runs": [P("ГАРАНТИРОВАН", 11, GOLD_LIGHT, F_DISPLAY, b=True, spc=450)],
           "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# СЛАЙД 3 — СКИДКА 10%
# ---------------------------------------------------------------------------
def slide_03(prs):
    slide = base_slide(prs, 3)
    ghost_number(slide, "02")
    kicker(slide, "ПРЕИМУЩЕСТВО 02 · ФИНАНСОВАЯ ОПТИМИЗАЦИЯ")
    heading(slide, "Скидка 10% на любую площадку")
    body(slide, "В разгар декабрьского ценового безумия мы за счет партнерских "
                "квот срезаем 10% со стоимости аренды абсолютно любой локации. "
                "Вы забираете статусное пространство на пике высокого сезона с "
                "очень приятной экономией для бюджета.")

    # Крупный акцентный бейдж «-10% НА АРЕНДУ»
    diamond(slide, 10.29, 2.16, 0.07)
    box(slide, MSO_SHAPE.OVAL, 8.69, 2.28, 3.2, 3.2,
        fill=GOLD, fill_a=4, line=GOLD, line_a=85, line_w=1.25)
    box(slide, MSO_SHAPE.OVAL, 9.09, 2.68, 2.4, 2.4,
        line=GOLD, line_a=35, line_w=0.75, dash=MSO_LINE.DASH)
    text(slide, 8.89, 3.06, 2.8, 0.3,
         [{"runs": [P("ПАРТНЕРСКАЯ КВОТА", 8, GRAY, F_DISPLAY, b=True, spc=350, a=80)],
           "align": PP_ALIGN.CENTER}])
    text(slide, 8.79, 3.36, 3.0, 0.95,
         [{"runs": [P("-10%", 54, GOLD_LIGHT, F_DISPLAY, b=True)],
           "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 8.79, 4.30, 3.0, 0.3,
         [{"runs": [P("НА АРЕНДУ", 10.5, GOLD, F_DISPLAY, b=True, spc=450)],
           "align": PP_ALIGN.CENTER}])
    text(slide, 8.15, 5.72, 4.28, 0.35,
         [{"runs": [P("Любая локация · Высокий сезон · Без ограничений", 8.5, GRAY,
                      F_BODY, spc=150, a=85)], "align": PP_ALIGN.CENTER}])


# ---------------------------------------------------------------------------
# СЛАЙД 4 — ТРИ КОНЦЕПЦИИ
# ---------------------------------------------------------------------------
def slide_04(prs):
    slide = base_slide(prs, 4)
    ghost_number(slide, "03")
    kicker(slide, "ПРЕИМУЩЕСТВО 03 · КРЕАТИВ И РЕЖИССУРА")
    heading(slide, "Написание 3-х вариантов подробной концепции мероприятия")
    body(slide, "Никаких банальных «голубых огоньков» — мы разработаем три "
                "уникальных авторских сценария, от иммерсивной зимней сказки до "
                "стильного светского бала в духе Met Gala. Вам останется лишь "
                "выбрать, какой вау-эффект ярче всего подчеркнет триумф компании "
                "в уходящем году.")

    text(slide, RX, 1.98, RW, 0.3,
         [{"runs": [P("ТРИ КОНЦЕПЦИИ · ОДИН ВАШ ВЫБОР", 8.5, GRAY,
                      F_DISPLAY, b=True, spc=250, a=80)]}])
    rows = [
        ("01", "Иммерсивное шоу", "зимняя сказка · интерактивный формат"),
        ("02", "Светский гала-ужин", "бальный вечер в духе Met Gala"),
        ("03", "Индивидуальный кастом", "полностью авторский сценарий"),
    ]
    for i, (num, title, sub) in enumerate(rows):
        y = 2.36 + i * 1.2
        box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, RX, y, RW, 1.02,
            fill=WHITE, fill_a=4, line=WHITE, line_a=10, line_w=0.75, radius=0.10)
        box(slide, MSO_SHAPE.RECTANGLE, RX, y + 0.26, 0.014, 0.5, fill=GOLD, fill_a=90)
        text(slide, 8.43, y, 0.5, 1.02,
             [{"runs": [P(num, 11, GOLD, F_DISPLAY, b=True, spc=100)]}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(slide, 8.95, y + 0.24, 3.3, 0.34,
             [{"runs": [P(title, 12.5, WHITE, F_DISPLAY, b=True)]}])
        text(slide, 8.95, y + 0.58, 3.3, 0.3,
             [{"runs": [P(sub, 9, GRAY, F_BODY, spc=50, a=85)]}])
    text(slide, RX, 6.02, RW, 0.3,
         [{"runs": [P("ПРЕЗЕНТАЦИЯ ТРЁХ ВАРИАНТОВ — ВХОДИТ В ПРЕДЗАКАЗ", 8, GOLD,
                      F_DISPLAY, b=True, spc=200, a=75)]}])


# ---------------------------------------------------------------------------
# СЛАЙД 5 — ПРЕДОПЛАТА 30%
# ---------------------------------------------------------------------------
def slide_05(prs):
    slide = base_slide(prs, 5)
    ghost_number(slide, "04")
    kicker(slide, "ПРЕИМУЩЕСТВО 04 · БЕЗОПАСНЫЙ СТАРТ")
    heading(slide, "Предоплата всего 30% от суммы мероприятия")
    body(slide, "Зафиксируйте за собой безупречный праздник всего за треть "
                "стоимости и спокойно готовьтесь к финалу года. Вы бронируете "
                "лучших подрядчиков и артистов прямо сейчас, не замораживая "
                "оборотные средства компании.")

    # Графическая шкала доли «30% для старта»
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, RX, 1.98, RW, 4.1,
        fill=WHITE, fill_a=4, line=WHITE, line_a=12, line_w=0.75, radius=0.045)
    box(slide, MSO_SHAPE.RECTANGLE, 8.45, 2.0, 3.68, 0.014, fill=GOLD, fill_a=80)
    text(slide, RXI, 2.34, RWI, 0.3,
         [{"runs": [P("СТРУКТУРА ФИКСАЦИИ БРОНИ", 8.5, GRAY, F_DISPLAY, b=True,
                      spc=300, a=80)]}])
    text(slide, 8.47, 2.66, 1.8, 1.15,
         [{"runs": [P("30%", 54, GOLD_LIGHT, F_DISPLAY, b=True)]}],
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 10.2, 2.66, 1.9, 1.15,
         [{"runs": [P("достаточно, чтобы дата и площадка стали вашими", 9.5,
                      GRAY, F_BODY, a=90)], "ls": 1.4}], anchor=MSO_ANCHOR.MIDDLE)
    # Шкала
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.47, 4.34, 3.64, 0.14,
        fill=WHITE, fill_a=9, radius=0.5)
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.47, 4.34, 1.09, 0.14,
        fill=GOLD, radius=0.5)
    line_v(slide, 9.555, 4.24, 0.34, GOLD, 90)
    text(slide, 8.47, 4.64, 0.3, 0.25, [{"runs": [P("0", 8, GRAY, F_BODY, a=70)]}])
    text(slide, 9.3, 4.64, 0.55, 0.25,
         [{"runs": [P("30", 8.5, GOLD, F_DISPLAY, b=True)], "align": PP_ALIGN.CENTER}])
    text(slide, 11.75, 4.64, 0.36, 0.25,
         [{"runs": [P("100", 8, GRAY, F_BODY, a=70)], "align": PP_ALIGN.RIGHT}])
    text(slide, RXI, 5.04, RWI, 0.62,
         [{"runs": [P("Остальные 70% — по комфортному графику после утверждения "
                      "концепции и подрядчиков.", 9.5, GRAY, F_BODY, a=90)],
           "ls": 1.4}])
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, RXI, 5.6, RWI, 0.38,
        line=GOLD, line_a=60, line_w=0.75, radius=0.5)
    text(slide, RXI, 5.6, RWI, 0.38,
         [{"runs": [P("БЕЗ ЗАМОРОЗКИ ОБОРОТНЫХ СРЕДСТВ", 8.5, GOLD_LIGHT,
                      F_DISPLAY, b=True, spc=150)], "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# СЛАЙД 6 — ГРАФИК ПЛАТЕЖЕЙ
# ---------------------------------------------------------------------------
def slide_06(prs):
    slide = base_slide(prs, 6)
    ghost_number(slide, "05")
    kicker(slide, "ПРЕИМУЩЕСТВО 05 · КОМФОРТ ДЛЯ БУХГАЛТЕРИИ")
    heading(slide, "Удобный график платежей")
    body(slide, "Никаких кассовых разрывов перед закрытием финансового года: "
                "мы составим комфортный календарь траншей, синхронизированный "
                "с вашей бухгалтерией. Вы закрываете бюджет поэтапно, "
                "прозрачно и без предновогоднего стресса.")

    # Линейный таймлайн: Транш 1 → Транш 2 → Финал
    text(slide, RX, 1.98, RW, 0.55,
         [{"runs": [P("КАЛЕНДАРЬ ТРАНШЕЙ · СОГЛАСУЕМ С ВАШЕЙ БУХГАЛЕРИЕЙ", 8.5,
                      GRAY, F_DISPLAY, b=True, spc=200, a=80)], "ls": 1.4}])
    box(slide, MSO_SHAPE.RECTANGLE, 8.45, 4.02, 3.68, 0.011, fill=WHITE, fill_a=16)
    nodes = [
        (8.6, "30%", "Транш 1", "фиксируем дату и площадку", "СЕГОДНЯ"),
        (10.29, "40%", "Транш 2", "сценарий и подрядчики", "НОЯБРЬ"),
        (11.98, "30%", "Финал", "полная готовность события", "ДЕКАБРЬ"),
    ]
    for cx, amt, name, det, when in nodes:
        box(slide, MSO_SHAPE.OVAL, cx - 0.14, 3.885, 0.28, 0.28,
            line=WHITE, line_a=30, line_w=0.75)
        box(slide, MSO_SHAPE.OVAL, cx - 0.065, 3.96, 0.13, 0.13, fill=GOLD)
        text(slide, cx - 0.75, 3.3, 1.5, 0.4,
             [{"runs": [P(amt, 15, GOLD_LIGHT, F_DISPLAY, b=True)],
               "align": PP_ALIGN.CENTER}])
        text(slide, cx - 0.85, 4.32, 1.7, 0.32,
             [{"runs": [P(name, 11.5, WHITE, F_DISPLAY, b=True)],
               "align": PP_ALIGN.CENTER}])
        text(slide, cx - 0.92, 4.68, 1.84, 0.55,
             [{"runs": [P(det, 8.5, GRAY, F_BODY, a=85)], "align": PP_ALIGN.CENTER,
               "ls": 1.3}])
        text(slide, cx - 0.85, 5.42, 1.7, 0.28,
             [{"runs": [P(when, 8, GOLD, F_DISPLAY, b=True, spc=350, a=80)],
               "align": PP_ALIGN.CENTER}])
    for ax in (9.4, 11.1):
        text(slide, ax, 3.92, 0.4, 0.3,
             [{"runs": [P("→", 11, GRAY, F_BODY, a=60)], "align": PP_ALIGN.CENTER}])


# ---------------------------------------------------------------------------
# СЛАЙД 7 — ПОДАРОК №1 (ТИМБИЛДИНГ)
# ---------------------------------------------------------------------------
def slide_07(prs):
    slide = base_slide(prs, 7)
    ghost_number(slide, "06")
    kicker(slide, "ПРЕИМУЩЕСТВО 06 · ПОДАРОК № 1 — ТИМБИЛДИНГ")
    heading(slide, "Предварительное онлайн-командообразующее мероприятие в подарок")
    body(slide, "Дарим интерактивный онлайн-квиз или предновогодний стрим, "
                "который снимет напряжение от декабрьских дедлайнов и объединит "
                "даже удаленщиков. Это идеальный разогрев, который подарит "
                "команде праздничное настроение задолго до первого звона бокалов.")

    # Плашка «SPECIAL GIFT / 0 ₽»
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, RX, 1.98, RW, 4.1,
        fill=GOLD, fill_a=4, line=GOLD, line_a=55, line_w=1.0, radius=0.045)
    box(slide, MSO_SHAPE.RECTANGLE, 8.45, 2.0, 3.68, 0.014, fill=GOLD, fill_a=90)
    text(slide, RXI, 2.42, RWI, 0.32,
         [{"runs": [P("SPECIAL GIFT", 10.5, GOLD_LIGHT, F_DISPLAY, b=True, spc=500)]}])
    text(slide, RXI, 2.84, RWI, 1.3,
         [{"runs": [P("0 ₽", 64, WHITE, F_DISPLAY, b=True)]}],
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, RXI, 4.24, RWI, 0.72,
         [{"runs": [P("Онлайн-квиз или предновогодний стрим — в подарок каждой "
                      "команде", 10, GRAY, F_BODY, a=90)], "ls": 1.4}])
    line_h(slide, RXI, 5.16, RWI, GOLD, 30)
    text(slide, RXI, 5.36, RWI, 0.5,
         [{"runs": [P("ИДЕАЛЬНЫЙ РАЗОГРЕВ · ДАЖЕ ДЛЯ УДАЛЁННЫХ КОМАНД", 8, GRAY,
                      F_DISPLAY, b=True, spc=250, a=75)], "ls": 1.4}])


# ---------------------------------------------------------------------------
# СЛАЙД 8 — ПОДАРОК №2 + CTA
# ---------------------------------------------------------------------------
def slide_08(prs):
    slide = base_slide(prs, 8)
    ghost_number(slide, "07")
    kicker(slide, "ПРЕИМУЩЕСТВО 07 · ПОДАРОК № 2 — ФИНАЛ")
    heading(slide, "Авторский шоу-торт с логотипом компании в подарок")
    body(slide, "Эффектная сладкая кульминация под бой курантов — за наш счет: "
                "создадим эксклюзивный торт с корпоративной символикой от "
                "звездного шеф-кондитера. Это не просто роскошный десерт за "
                "успехи года, а главный герой праздничных сторис ваших "
                "сотрудников!")

    # Мини-карточка финала
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, RX, 1.98, RW, 2.9,
        fill=GOLD, fill_a=4, line=GOLD, line_a=55, line_w=1.0, radius=0.045)
    box(slide, MSO_SHAPE.RECTANGLE, 8.45, 2.0, 3.68, 0.014, fill=GOLD, fill_a=90)
    text(slide, RXI, 2.3, RWI, 0.3,
         [{"runs": [P("FINALE · 0 ₽", 9.5, GOLD_LIGHT, F_DISPLAY, b=True, spc=400)]}])
    text(slide, RXI, 2.58, RWI, 0.95,
         [{"runs": [P("0 ₽", 44, WHITE, F_DISPLAY, b=True)]}])
    text(slide, RXI, 3.6, RWI, 0.62,
         [{"runs": [P("Авторский шоу-торт от звездного шеф-кондитера — "
                      "с логотипом компании", 9.5, GRAY, F_BODY, a=90)], "ls": 1.4}])
    line_h(slide, RXI, 4.3, RWI, GOLD, 30)
    text(slide, RXI, 4.42, RWI, 0.3,
         [{"runs": [P("ГЛАВНЫЙ ГЕРОЙ ПРАЗДНИЧНЫХ СТОРИС", 8, GRAY, F_DISPLAY,
                      b=True, spc=150, a=70)]}])

    # CTA-блок
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.42, 11.53, 1.28,
        fill=WHITE, fill_a=4, line=GOLD, line_a=55, line_w=1.0, radius=0.09)
    diamond(slide, 1.295, 6.06, 0.09)
    text(slide, 1.5, 5.42, 5.5, 1.28,
         [{"runs": [P("ЗАБРОНИРУЙТЕ ВСТРЕЧУ И ПРЕЗЕНТАЦИЮ ТРЁХ КОНЦЕПЦИЙ", 13.5,
                      WHITE, F_DISPLAY, b=True, spc=60)], "ls": 1.2},
          {"runs": [P("УЖЕ НА ЭТОЙ НЕДЕЛЕ", 9.5, GOLD_LIGHT, F_DISPLAY, b=True,
                      spc=300)], "before": 4}],
         anchor=MSO_ANCHOR.MIDDLE)
    line_v(slide, 7.1, 5.72, 0.68, GOLD, 30)
    text(slide, 7.3, 5.42, 4.85, 1.28,
         [{"runs": [P("ТЕЛЕФОН · +7 900 000-00-00", 9, GRAY, F_BODY, spc=120, a=95)],
           "align": PP_ALIGN.RIGHT, "ls": 1.55},
          {"runs": [P("ПОЧТА · HELLO@AGENCY.RU", 9, GRAY, F_BODY, spc=120, a=95)],
           "align": PP_ALIGN.RIGHT, "ls": 1.55},
          {"runs": [P("TELEGRAM · @AGENCY", 9, GRAY, F_BODY, spc=120, a=95)],
           "align": PP_ALIGN.RIGHT, "ls": 1.55}],
         anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# СБОРКА
# ---------------------------------------------------------------------------
def build(out_pptx, out_scene):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Предварительный заказ новогоднего корпоратива"
    prs.core_properties.author = "Event Agency · Moscow"
    prs.core_properties.subject = "Ключевые преимущества предзаказа · Сезон 2024–2025"

    slide_01(prs)
    slide_02(prs)
    slide_03(prs)
    slide_04(prs)
    slide_05(prs)
    slide_06(prs)
    slide_07(prs)
    slide_08(prs)

    prs.save(out_pptx)
    with open(out_scene, "w", encoding="utf-8") as fh:
        json.dump(SCENE, fh, ensure_ascii=False, indent=1)
    print("OK  ->", out_pptx)
    print("OK  ->", out_scene)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    build(os.path.join(here, "preorder_new_year_corporate_2025.pptx"),
          os.path.join(here, "scene.json"))
